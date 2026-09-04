#!/usr/bin/env python3
"""Раскладка «услуга × город» и сборка фида поверх текущего.

Правила раскладки (см. README, раздел про антидубль 2015):
  * 19 услуг стоят в каждом городе — их объекты работы не пересекаются;
  * три группы «ведение учёта» конфликтуют внутри себя, поэтому в городе
    стоит РОВНО ОДИН представитель каждой группы, с перевесом на самый
    универсальный вариант;
  * пара «услуга × город», уже занятая живым объявлением, пропускается —
    иначе новая строка станет точным дублем живой.

    python plan_build.py <текущий_фид.xlsx> <feed.xlsx> [--csv матрица.csv]
"""
import argparse, csv, json, os, sys

import openpyxl

sys.path.insert(0, os.environ.get(
    'AVITO_SKILL', '/root/.claude/skills/synced/'
    '3a0975c2-241e-420c-9584-bd7bc9dba0fe_113d8e35-a8bd-421c-bcf2-b9d9e9f4dc01/avito-manager/scripts'))
import build_feed

from build_city_feed import make_description, DISK_DIR, COMMON_PHOTOS
from cities import CITIES

SHEET = "Деловые услуги-Бухгалтерия, фин"
ID_PREFIX = "260904"
# у строки-образца Price пустой, объявление ушло бы без цены; 1000 ₽ — то,
# что стояло у большинства объявлений этого аккаунта раньше
NEW_DEFAULTS = {"Price": "1000"}

# группы «ведение учёта»: в городе ровно один представитель, первый — основной
ROTATIONS = [
    ["01", "02", "03", "10"],   # стройка
    ["13", "11", "12", "20"],   # маркетплейсы
    ["21", "22", "23"],         # транспорт
]


def rotation_plan(members, n_cities, offset):
    """Основной вариант в половине городов, остальные делят вторую половину."""
    major, minors = members[0], members[1:]
    out = []
    for i in range(n_cities):
        j = (i + offset) % n_cities
        out.append(major if j % 2 == 0 else minors[(j // 2) % len(minors)])
    return out


def load_ads():
    raw = json.load(open(os.path.join("data", "ads_texts.json"), encoding="utf-8"))
    ads = {}
    for num, key in enumerate(sorted(raw), 1):
        lines = raw[key].strip().split("\n")
        title = lines[0].strip()
        body = "\n".join(lines[1:]).strip()
        ads["%02d" % num] = {
            "num": "%02d" % num, "title": title,
            "desc": make_description(title, body),
            "cover": "%s/cover_%02d.jpg" % (DISK_DIR, num),
        }
    return ads


def make_base(current, active_ids, base_path):
    """Оставляет в фиде только живые строки.

    Всё остальное в скачанном файле — снятые объявления. Если оставить их,
    выгрузка попробует опубликовать архив заново и сожжёт слоты пакета.
    Возвращает пары (заголовок, адрес) живых — их новыми не дублируем.
    """
    wb = openpyxl.load_workbook(current)
    live = set()
    for name in wb.sheetnames:
        if name.startswith("Спр-") or name == "Инструкция":
            continue
        ws = wb[name]
        h = [str(ws.cell(2, c).value or "").strip() for c in range(1, ws.max_column + 1)]
        col = {n: i + 1 for i, n in enumerate(h) if n}
        if "Id" not in col:
            continue
        drop = []
        for r in range(5, ws.max_row + 1):
            if not ws.cell(r, col["Id"]).value:
                continue
            if str(ws.cell(r, col["AvitoId"]).value).strip() in active_ids:
                live.add((str(ws.cell(r, col["Title"]).value).strip(),
                          str(ws.cell(r, col["Address"]).value).strip()))
            else:
                drop.append(r)
        for r in reversed(drop):
            ws.delete_rows(r)
        if drop:
            print("лист %r: убрано снятых строк %d" % (name, len(drop)))
    wb.save(base_path)
    return live


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("current"); ap.add_argument("out")
    ap.add_argument("--csv", default="reports/matrix.csv")
    ap.add_argument("--recon", default="snapshots/recon_2026-09-04.json",
                    help="снимок разведки: из него берём AvitoId живых объявлений")
    ap.add_argument("--base", default="reports/base.xlsx")
    a = ap.parse_args()

    ads = load_ads()
    rotating = {n for grp in ROTATIONS for n in grp}
    always = [n for n in sorted(ads) if n not in rotating]
    print("услуг всего %d: в каждом городе %d, в ротациях %d"
          % (len(ads), len(always), len(rotating)))

    plans = [rotation_plan(g, len(CITIES), off) for off, g in enumerate(ROTATIONS)]
    active_ids = {str(i["id"]) for i in json.load(
        open(a.recon, encoding="utf-8"))["active_items"]}
    os.makedirs(os.path.dirname(a.base) or ".", exist_ok=True)
    occupied = make_base(a.current, active_ids, a.base)
    print("живых объявлений: %d, их пары «услуга × город» пропускаем" % len(occupied))

    # живое объявление занимает не только свою пару, но и всю свою группу:
    # второй представитель конфликтной группы в этом городе даст 2015
    by_title = {ad["title"]: num for num, ad in ads.items()}
    live_groups = {}
    for title, address in occupied:
        num = by_title.get(title)
        for gi, grp in enumerate(ROTATIONS):
            if num in grp:
                live_groups.setdefault(address, set()).add(gi)

    items, matrix, skipped = [], [], []
    for i, city in enumerate(CITIES):
        address = city[1]
        slots = list(always)
        for gi, p in enumerate(plans):
            if gi in live_groups.get(address, ()):
                skipped.append(("группа %d" % (gi + 1), city[0], "занята живым"))
                continue
            slots.append(p[i])
        for num in slots:
            ad = ads[num]
            if (ad["title"], address) in occupied:
                skipped.append((num, city[0], "точный дубль живого"))
                continue
            items.append({"sheet": SHEET, "title": ad["title"], "desc": ad["desc"],
                          "addr": address, "cover": ad["cover"], "avito_id": None})
            matrix.append((city[0], address, num, ad["title"]))
    print("пропущено как уже живое: %s" % (skipped or "нет"))
    print("новых строк: %d" % len(items))

    tail = COMMON_PHOTOS
    added, cleared = build_feed.build(a.base, items, a.out, ID_PREFIX,
                                      tails={SHEET: tail})
    print("добавлено: %s | очищено служебных ячеек: %d" % (added, cleared))

    wb = openpyxl.load_workbook(a.out)
    ws = wb[SHEET]
    h = [str(ws.cell(2, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    col = {n: i + 1 for i, n in enumerate(h) if n}
    filled = 0
    for r in range(5, ws.max_row + 1):
        rid = str(ws.cell(r, col["Id"]).value or "")
        if not rid.startswith(ID_PREFIX):      # живые строки не трогаем
            continue
        for name, val in NEW_DEFAULTS.items():
            if ws.cell(r, col[name]).value in (None, ""):
                ws.cell(r, col[name]).value = val
                filled += 1
    wb.save(a.out)
    print("проставлено значений по умолчанию у новых строк: %d" % filled)

    os.makedirs(os.path.dirname(a.csv) or ".", exist_ok=True)
    with open(a.csv, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["город", "адрес", "№", "объявление"])
        w.writerows(matrix)
    print("матрица: %s" % a.csv)


if __name__ == "__main__":
    main()
