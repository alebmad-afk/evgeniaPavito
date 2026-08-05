#!/usr/bin/env python3
"""Сборка feed.xlsx для автовыгрузки Авито из официального Excel-экспорта.

Оставляет в файле только объявления с действием KEEP/ACTIVATE из отчёта
(reports/report_*.csv); строки DELETE удаляются — при загрузке такого файла
Авито снимает отсутствующие объявления с публикации, а присутствующие
публикует/активирует.

Использование:
    python3 build_feed_xlsx.py <экспорт_авито.xlsx> <отчёт.csv> <feed.xlsx>
"""
import csv
import sys

import openpyxl

SERVICE_SHEETS_SKIP = ("Инструкция", "Спр-")


def main():
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    src, report, dst = sys.argv[1], sys.argv[2], sys.argv[3]

    action = {}
    with open(report, encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            action[r["item_id"]] = r["planned_action"]
    keep_ids = {i for i, a in action.items() if a in ("KEEP", "ACTIVATE")}
    print(f"в отчёте KEEP/ACTIVATE: {len(keep_ids)}")

    wb = openpyxl.load_workbook(src)
    kept_avito_ids = set()
    for name in wb.sheetnames:
        if name.startswith(SERVICE_SHEETS_SKIP):
            continue
        ws = wb[name]
        header = [c.value for c in ws[2]]
        i_avito = header.index("AvitoId")
        to_delete = []
        kept = 0
        for ri in range(5, ws.max_row + 1):
            rid = ws.cell(row=ri, column=1).value
            if rid is None or str(rid).strip() == "":
                continue
            avito_id = str(ws.cell(row=ri, column=i_avito + 1).value or "").strip()
            if avito_id in keep_ids:
                kept += 1
                kept_avito_ids.add(avito_id)
            else:
                to_delete.append(ri)
        for ri in reversed(to_delete):
            ws.delete_rows(ri)
        print(f"{name}: оставлено {kept}, удалено {len(to_delete)}")

    wb.save(dst)

    # валидация результата
    wb2 = openpyxl.load_workbook(dst)
    total = 0
    check_ids = set()
    for name in wb2.sheetnames:
        if name.startswith(SERVICE_SHEETS_SKIP):
            continue
        ws = wb2[name]
        header = [c.value for c in ws[2]]
        i_avito = header.index("AvitoId")
        for row in ws.iter_rows(min_row=5, values_only=True):
            if row[0] is None or str(row[0]).strip() == "":
                continue
            total += 1
            check_ids.add(str(row[i_avito] or "").strip())
    lost = kept_avito_ids - check_ids
    extra = check_ids - keep_ids
    print(f"итог: {total} объявлений в {dst}; потеряно: {len(lost)}; лишних: {len(extra)}")
    if lost or extra:
        sys.exit("ОШИБКА: несоответствие после фильтрации")
    missing = keep_ids - check_ids
    print(f"KEEP/ACTIVATE, которых нет в файле (активировать вручную): {len(missing)}")
    for m in sorted(missing):
        print("  ", m, action[m])


if __name__ == "__main__":
    main()
